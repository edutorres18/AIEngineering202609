"""Adjuntos por el camino B (extracción local): el texto de PDF y Word se saca aquí y viaja al
LLM como parte del mensaje del usuario.

Frente al camino A (subir el binario a la Files API del proveedor), el B no ata el servicio a
un proveedor multimodal: el wrapper sigue siendo «texto entra, texto sale» con OpenAI y con
Anthropic. Se pierde lo visual (diagramas, imágenes), pero se controla qué entra en el prompt y
prepara el terreno para el chunking de RAG del módulo 3.
"""

import io
import logging
from pathlib import PurePath

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".pdf", ".docx")


class AttachmentError(Exception):
    """No se pudo leer un adjunto (formato no admitido o archivo dañado)."""

    def __init__(self, filename: str, message: str, *, status_code: int = 422) -> None:
        super().__init__(message)
        self.filename = filename
        self.message = message
        self.status_code = status_code


def _pdf_text(content: bytes) -> str:
    from pypdf import PdfReader

    pages = []
    for number, page in enumerate(PdfReader(io.BytesIO(content)).pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # una página dañada no invalida el resto del documento
            logger.warning("No se pudo extraer la página %d del PDF", number)
    return "\n\n".join(page for page in pages if page.strip())


def _docx_text(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:  # las especificaciones suelen llevar tablas
        lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(line for line in lines if line.strip())


def extract_text(filename: str, content: bytes, *, max_chars: int) -> str:
    """Texto del adjunto, recortado a `max_chars`. Vacío si el documento no tiene texto
    (p. ej. un PDF escaneado: sin OCR no hay nada que extraer)."""
    extension = PurePath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise AttachmentError(
            filename,
            f"«{filename}» no es un formato admitido: solo PDF (.pdf) y Word (.docx).",
            status_code=415,
        )
    try:
        text = (_pdf_text if extension == ".pdf" else _docx_text)(content).strip()
    except Exception as exc:
        logger.warning("No se pudo leer el adjunto %s: %s", filename, type(exc).__name__)
        raise AttachmentError(
            filename, f"No se pudo leer «{filename}»: el archivo parece dañado."
        ) from exc
    if len(text) > max_chars:
        logger.info("Adjunto %s recortado de %d a %d caracteres", filename, len(text), max_chars)
        text = text[:max_chars]
    return text


def enrich_transcript(transcript: str, attachments: list[tuple[str, str]]) -> str:
    """Transcripción seguida del texto de cada adjunto, entre separadores con su nombre."""
    parts = [transcript.strip()]
    for filename, text in attachments:
        if text:
            parts.append(f"--- attachment: {filename} ---\n{text}\n--- end attachment ---")
    return "\n\n".join(parts)
