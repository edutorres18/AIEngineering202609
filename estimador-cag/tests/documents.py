"""Documentos de prueba generados en memoria: no hace falta guardar binarios en el repo."""

import io


def pdf_bytes(*lines: str) -> bytes:
    """PDF mínimo de una página con esas líneas de texto (ASCII, sin paréntesis)."""
    content = "BT /F1 12 Tf 72 720 Td 16 TL " + " ".join(f"({line}) Tj T*" for line in lines)
    content += " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    pdf = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    pdf += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    pdf += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    pdf += f"startxref\n{xref}\n%%EOF\n".encode()
    return pdf


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """Documento Word con esos párrafos y, opcionalmente, una tabla."""
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(word_table.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
